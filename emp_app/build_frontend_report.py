from docx import Document
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_CELL_VERTICAL_ALIGNMENT
from docx.enum.style import WD_STYLE_TYPE
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.enum.section import WD_SECTION
from pathlib import Path

OUT = Path(r"D:\emp_app\PulseHR_Frontend_Implementation_Report.docx")

BLUE = RGBColor(46, 116, 181)
DARK_BLUE = RGBColor(31, 77, 120)
MUTED = RGBColor(89, 99, 112)
INK = RGBColor(28, 42, 60)
HEADER_FILL = "F2F4F7"
CONTENT_WIDTH = 9360

def set_cell_shading(cell, fill):
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = OxmlElement('w:shd')
    shd.set(qn('w:fill'), fill)
    tc_pr.append(shd)

def set_cell_width(cell, width):
    tc_pr = cell._tc.get_or_add_tcPr()
    tc_w = tc_pr.find(qn('w:tcW'))
    if tc_w is None:
        tc_w = OxmlElement('w:tcW')
        tc_pr.append(tc_w)
    tc_w.set(qn('w:w'), str(width))
    tc_w.set(qn('w:type'), 'dxa')

def set_table_geometry(table, widths):
    table.autofit = False
    table.alignment = WD_TABLE_ALIGNMENT.LEFT
    tbl_pr = table._tbl.tblPr
    tbl_w = tbl_pr.first_child_found_in('w:tblW')
    if tbl_w is None:
        tbl_w = OxmlElement('w:tblW')
        tbl_pr.append(tbl_w)
    tbl_w.set(qn('w:w'), str(sum(widths)))
    tbl_w.set(qn('w:type'), 'dxa')
    tbl_ind = OxmlElement('w:tblInd')
    tbl_ind.set(qn('w:w'), '120')
    tbl_ind.set(qn('w:type'), 'dxa')
    tbl_pr.append(tbl_ind)
    grid = table._tbl.tblGrid
    for col, width in zip(grid.gridCol_lst, widths):
        col.set(qn('w:w'), str(width))
    for row in table.rows:
        for cell, width in zip(row.cells, widths):
            set_cell_width(cell, width)
            cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
            tc_pr = cell._tc.get_or_add_tcPr()
            margins = OxmlElement('w:tcMar')
            for side, value in [('top', '80'), ('bottom', '80'), ('start', '120'), ('end', '120')]:
                node = OxmlElement(f'w:{side}')
                node.set(qn('w:w'), value)
                node.set(qn('w:type'), 'dxa')
                margins.append(node)
            tc_pr.append(margins)

def set_font(run, size=11, bold=None, color=INK):
    run.font.name = 'Calibri'
    run._element.rPr.rFonts.set(qn('w:ascii'), 'Calibri')
    run._element.rPr.rFonts.set(qn('w:hAnsi'), 'Calibri')
    run.font.size = Pt(size)
    run.font.color.rgb = color
    if bold is not None:
        run.bold = bold

def add_text(doc, text, style='Normal', after=None):
    p = doc.add_paragraph(style=style)
    r = p.add_run(text)
    set_font(r)
    if after is not None:
        p.paragraph_format.space_after = Pt(after)
    return p

def add_bullets(doc, items):
    for item in items:
        p = doc.add_paragraph(style='List Bullet')
        p.paragraph_format.space_after = Pt(4)
        p.paragraph_format.line_spacing = 1.167
        set_font(p.add_run(item))

def add_heading(doc, text, level=1):
    p = doc.add_paragraph(style=f'Heading {level}')
    set_font(p.add_run(text), size=16 if level == 1 else 13, bold=True, color=BLUE if level < 3 else DARK_BLUE)
    return p

def add_table(doc, headers, rows, widths):
    table = doc.add_table(rows=1, cols=len(headers))
    table.style = 'Table Grid'
    set_table_geometry(table, widths)
    for cell, label in zip(table.rows[0].cells, headers):
        set_cell_shading(cell, HEADER_FILL)
        p = cell.paragraphs[0]
        p.paragraph_format.space_after = Pt(0)
        set_font(p.add_run(label), size=10, bold=True, color=INK)
    for row_values in rows:
        cells = table.add_row().cells
        for cell, value in zip(cells, row_values):
            p = cell.paragraphs[0]
            p.paragraph_format.space_after = Pt(0)
            set_font(p.add_run(value), size=10)
    return table

doc = Document()
section = doc.sections[0]
section.top_margin = Inches(1)
section.bottom_margin = Inches(1)
section.left_margin = Inches(1)
section.right_margin = Inches(1)
section.header_distance = Inches(0.492)
section.footer_distance = Inches(0.492)

styles = doc.styles
normal = styles['Normal']
normal.font.name = 'Calibri'
normal._element.rPr.rFonts.set(qn('w:ascii'), 'Calibri')
normal._element.rPr.rFonts.set(qn('w:hAnsi'), 'Calibri')
normal.font.size = Pt(11)
normal.paragraph_format.space_after = Pt(6)
normal.paragraph_format.line_spacing = 1.10
for name, size, before, after, color in [('Heading 1', 16, 16, 8, BLUE), ('Heading 2', 13, 12, 6, BLUE), ('Heading 3', 12, 8, 4, DARK_BLUE)]:
    st = styles[name]
    st.font.name = 'Calibri'
    st._element.rPr.rFonts.set(qn('w:ascii'), 'Calibri')
    st._element.rPr.rFonts.set(qn('w:hAnsi'), 'Calibri')
    st.font.size = Pt(size)
    st.font.bold = True
    st.font.color.rgb = color
    st.paragraph_format.space_before = Pt(before)
    st.paragraph_format.space_after = Pt(after)

# Quiet header/footer
header = section.header.paragraphs[0]
header.alignment = WD_ALIGN_PARAGRAPH.LEFT
set_font(header.add_run('Pulse HRMS | Frontend Implementation Report'), size=9, color=MUTED)
footer = section.footer.paragraphs[0]
footer.alignment = WD_ALIGN_PARAGRAPH.RIGHT
set_font(footer.add_run('Frontend-only delivery | 13 August 2026'), size=9, color=MUTED)

# Memo masthead
p = doc.add_paragraph()
p.paragraph_format.space_after = Pt(4)
r = p.add_run('FRONTEND IMPLEMENTATION REPORT')
set_font(r, size=23, bold=True, color=INK)
p = doc.add_paragraph()
p.paragraph_format.space_after = Pt(14)
set_font(p.add_run('Pulse HRMS Mobile Application - Three-Day Delivery Summary'), size=14, color=MUTED)

meta = [('Project', 'Pulse HRMS Flutter frontend'), ('Delivery period', 'Three implementation days'), ('Status', 'Frontend completed with local mock data'), ('Backend scope', 'Deferred for the next phase')]
for label, value in meta:
    p = doc.add_paragraph()
    p.paragraph_format.space_after = Pt(2)
    set_font(p.add_run(f'{label}: '), bold=True)
    set_font(p.add_run(value))

add_heading(doc, 'Executive Summary')
add_text(doc, 'This report records the frontend completion work carried out for the Pulse HRMS Flutter application. The goal was to bring the Flutter project to functional parity with the supplied PulseHR mobile reference while keeping APIs, database integration, and production services out of scope. The result is a frontend-only application with local mock data, connected navigation, user-facing forms, details, confirmations, role-aware access, and a clean backend handoff boundary.')

add_heading(doc, 'Scope Delivered')
add_bullets(doc, [
    'Completed employee and manager interface flows using local state and mock data.',
    'Connected home cards, quick actions, notifications, detail pages, forms, modal sheets, and sign-out behavior.',
    'Implemented role-aware navigation so employee and manager experiences remain distinct during a session.',
    'Kept service classes intentionally unimplemented because backend and API work are planned for the next phase.',
])

add_heading(doc, 'Day 1 - Review, Foundation, and Core Flows')
add_text(doc, 'The first day focused on comparing the Flutter project with the reference application, identifying incomplete frontend areas, and establishing the base application flow.')
add_bullets(doc, [
    'Reviewed the reference implementation and the Flutter screens to identify missing details, routes, and interaction states.',
    'Established the application shell with Home, Attendance, Leave, Notifications, and More modules.',
    'Completed the main attendance experience: clock in/out state, running timer, history filters, record details, calendar navigation, and attendance correction submission feedback.',
    'Completed the leave experience: balances, history, leave application form, date selection, optional attachment state, request details, cancellation path, and manager review actions.',
    'Added the initial mock-data approach so every frontend action could be demonstrated without backend dependencies.',
])

add_heading(doc, 'Day 2 - Module Completion and Shared UI')
add_text(doc, 'The second day expanded the implementation across the remaining employee modules and removed screen/widget placeholders that would otherwise create dead-end navigation.')
add_table(doc, ['Area', 'Frontend behavior completed'], [
    ('Authentication', 'Sign in, sign-up validation, department and role selection, remember-me control, terms confirmation, password visibility, and password-reset feedback.'),
    ('Team and performance', 'Team directory, search/filter behavior, member profile details, team attendance, team performance KPIs, goals, and review details.'),
    ('Documents and payslips', 'Document categories, detail sheets, PDF-style preview flow, download feedback, payslip history, earnings/deductions summary, and payslip preview flow.'),
    ('Helpdesk and requests', 'Ticket category/priority selection, ticket replies, local ticket creation, self-service request creation, details, and local status display.'),
    ('Calendar and announcements', 'Calendar selection and filters, event list, announcement read state, full announcement detail, and home navigation paths.'),
], [1900, 7460])

add_heading(doc, 'Day 3 - Role Access, Session Behavior, and Quality Checks')
add_text(doc, 'The third day focused on correct employee/manager separation, session actions, and final verification of the Flutter frontend.')
add_bullets(doc, [
    'Separated employee and manager navigation. Employee navigation uses Home, Attendance, Leave, Alerts, and More; manager navigation uses Home, Attendance, My Team, Leave, and More.',
    'Restricted manager-only tools: Leave Approvals, My Team, Team Attendance, and Team Performance are visible and callable only in manager mode.',
    'Added route guarding so manager-only flows cannot be opened from employee mode through future accidental links.',
    'Changed the role label in the app bar to display-only. The role no longer changes by clicking the label; the session remains tied to the authenticated user role.',
    'Made manager selection during sign-up initialize the manager session. Employee sign-up initializes employee mode.',
    'Implemented real frontend sign-out from both Profile and More, returning the user to the login screen.',
    'Expanded settings with notification controls, theme, biometric UI state, device frame selector, simulated app states, reset-demo feedback, and frontend implementation notes.',
])

add_heading(doc, 'Role and Navigation Rules')
add_table(doc, ['User type', 'Available navigation and tools'], [
    ('Employee', 'Home, Attendance, Leave, Alerts, More, Payslips, Performance & Goals, Helpdesk, Calendar, Self-Service Requests, Documents, Profile, and Settings.'),
    ('Manager', 'Home, Attendance, My Team, Leave, More, plus manager tools for Leave Approvals, Team Attendance, and Team Performance. Employee services remain available where applicable.'),
], [1700, 7660])

add_heading(doc, 'Verification Performed')
add_bullets(doc, [
    'Formatted the modified Dart files using the Dart formatter.',
    'Ran Dart static analysis on the Flutter source: passed without reported issues.',
    'Ran Flutter widget tests: passed.',
    'Built a debug Android APK successfully during the completion pass.',
])

add_heading(doc, 'Backend Handoff / Next Phase')
add_text(doc, 'The delivered application is intentionally frontend-only. Local mock data and local state should be replaced by repositories and service implementations in the backend phase. The next phase should connect authentication, employee profile, attendance, leave, approvals, notifications, helpdesk tickets, documents, payslips, requests, and role claims to real API and database services. The UI contracts and role rules implemented in this delivery provide the structure for that integration.')

add_heading(doc, 'Acceptance Summary')
add_text(doc, 'The Pulse HRMS frontend is complete for the agreed mock-data scope. Users can navigate through the implemented employee or manager experience, perform the designed frontend interactions, receive local UI feedback, sign out, and return to login. Backend data persistence and external service connectivity remain the explicit next step.')

doc.core_properties.title = 'Pulse HRMS Frontend Implementation Report'
doc.core_properties.subject = 'Three-Day Frontend Delivery Summary'
doc.core_properties.author = 'Pulse HRMS Frontend Team'
doc.save(OUT)
print(OUT)
